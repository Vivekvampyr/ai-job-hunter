import os
from typing import Tuple
from pypdf import PdfReader
from docx import Document


class ResumeParserService:
    """Extracts raw text and metadata from PDF and DOCX resume files."""

    @staticmethod
    def parse_file(file_path: str) -> Tuple[str, str, int]:
        """Extracts text, file type, and file size from a resume file."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Resume file not found at {file_path}")

        file_size = os.path.getsize(file_path)
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            raw_text = ResumeParserService._parse_pdf(file_path)
            file_type = "pdf"
        elif ext in [".docx", ".doc"]:
            raw_text = ResumeParserService._parse_docx(file_path)
            file_type = "docx"
        else:
            raise ValueError(f"Unsupported resume file format '{ext}'. Allowed: .pdf, .docx")

        clean_text = raw_text.strip()
        return clean_text, file_type, file_size

    @staticmethod
    def _parse_pdf(file_path: str) -> str:
        text_parts = []
        with open(file_path, "rb") as f:
            reader = PdfReader(f)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_parts.append(extracted)
        return "\n".join(text_parts)

    @staticmethod
    def _parse_docx(file_path: str) -> str:
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also include table text if any
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)
