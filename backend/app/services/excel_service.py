import io
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelExportService:
    """Exports structured job listings with contacts, match scores, and application links to Excel."""

    COLUMNS = [
        ("Company", 24),
        ("Role", 32),
        ("Location", 24),
        ("Work Mode", 14),
        ("Match", 14),
        ("Email", 28),
        ("Phone", 20),
        ("Application Form", 35),
        ("Apply URL", 35)
    ]

    @classmethod
    def generate_excel(cls, jobs_data: List[Dict[str, Any]]) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "AI Job Hunter Matches"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10, color="0F172A")
        link_font = Font(name="Calibri", size=10, color="2563EB", underline="single")
        
        thin_border_side = Side(border_style="thin", color="E2E8F0")
        row_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        
        # Match Level Badges
        high_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        high_font = Font(name="Calibri", size=10, bold=True, color="166534")

        med_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        med_font = Font(name="Calibri", size=10, bold=True, color="92400E")

        low_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        low_font = Font(name="Calibri", size=10, color="475569")

        # Write Header Row
        for col_idx, (col_name, _) in enumerate(cls.COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
            cell.border = row_border
        ws.row_dimensions[1].height = 28

        # Write Data Rows
        for row_idx, job in enumerate(jobs_data, start=2):
            match_level = job.get("match_level", "Medium")
            apply_url = job.get("apply_url", "")
            app_form = job.get("application_form", "") or apply_url

            row_values = [
                job.get("company_name", ""),
                job.get("title", ""),
                job.get("location", ""),
                job.get("work_mode", "Remote"),
                f"{match_level} ({job.get('match_score', 50)}%)",
                job.get("email", "Not Found"),
                job.get("phone", "Not Found"),
                app_form,
                apply_url
            ]

            for col_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = row_border
                cell.alignment = Alignment(vertical="center")

                # Center align location, work mode
                if col_idx in [3, 4]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Match column special formatting
                if col_idx == 5:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if match_level == "High":
                        cell.fill = high_fill
                        cell.font = high_font
                    elif match_level == "Medium":
                        cell.fill = med_fill
                        cell.font = med_font
                    else:
                        cell.fill = low_fill
                        cell.font = low_font

                # Hyperlinks for URLs
                if col_idx in [8, 9] and str(val).startswith("http"):
                    cell.hyperlink = val
                    cell.font = link_font

            ws.row_dimensions[row_idx].height = 22

        # Set Column Widths
        for col_idx, (_, width) in enumerate(cls.COLUMNS, start=1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
